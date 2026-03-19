#!/usr/bin/env python3
"""提取 Word 文档内容"""

import docx
import sys

def extract_text(docx_path):
    """提取 Word 文档中的文字"""
    try:
        doc = docx.Document(docx_path)
        text = []
        
        # 提取段落
        for para in doc.paragraphs:
            if para.text.strip():
                text.append(para.text)
        
        # 提取表格
        for table in doc.tables:
            text.append("\n[表格]")
            for row in table.rows:
                cells = [cell.text for cell in row.cells]
                text.append(" | ".join(cells))
            text.append("")
        
        return "\n".join(text)
    
    except Exception as e:
        return f"错误：{e}"

if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("用法：extract_docx.py <文件路径>")
        sys.exit(1)
    
    content = extract_text(sys.argv[1])
    print(content)
